"""Tests for Phase 7A: Multi-format document parser (PDF, DOCX, TXT)."""

from __future__ import annotations

import io
from pathlib import Path
import pytest
import docx
import fitz

from app.parsers.document_parser import (
    DocumentParsingError,
    extract_text_from_document,
    extract_text_from_docx,
    extract_text_from_txt,
)


@pytest.fixture
def sample_docx_bytes():
    """Create in-memory DOCX bytes with sample candidate resume text."""
    doc = docx.Document()
    doc.add_heading("Marcus Vance", 0)
    doc.add_paragraph("marcus.vance@example.com | Austin, TX")
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Node.js, Express, MongoDB, REST API, Git, Docker")
    doc.add_heading("Experience", level=1)
    p = doc.add_paragraph("Backend Developer Intern at DataPulse")
    doc.add_paragraph("Designed secure REST APIs with Node.js and Express.")
    
    # Add a table
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Degree"
    hdr_cells[1].text = "B.S. Software Engineering"

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture
def sample_pdf_bytes():
    """Create in-memory PDF bytes with sample text."""
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(54, 54, 500, 500)
    page.insert_textbox(rect, "Rahul Sharma\nReact, Node.js, MongoDB, REST API, Git\nFull Stack Developer", fontsize=11)
    stream = io.BytesIO()
    doc.save(stream)
    doc.close()
    return stream.getvalue()


@pytest.fixture
def sample_txt_bytes():
    """Create in-memory UTF-8 text bytes."""
    text = "Aisha Khan\nSkills: React, Node.js, Express, MongoDB, AWS, Git\nEducation: B.S. CS"
    return text.encode("utf-8")


class TestDocumentParsers:
    """Test suite for individual document formats."""

    def test_extract_from_docx_bytes(self, sample_docx_bytes):
        text = extract_text_from_docx(sample_docx_bytes)
        assert "Marcus Vance" in text
        assert "Node.js" in text
        assert "DataPulse" in text
        assert "B.S. Software Engineering" in text

    def test_extract_from_docx_file(self, tmp_path, sample_docx_bytes):
        path = tmp_path / "resume.docx"
        path.write_bytes(sample_docx_bytes)
        text = extract_text_from_docx(path)
        assert "Marcus Vance" in text

    def test_extract_from_empty_docx_raises(self, tmp_path):
        path = tmp_path / "empty.docx"
        path.write_bytes(b"")
        with pytest.raises(DocumentParsingError, match="empty"):
            extract_text_from_docx(path)

    def test_extract_from_corrupted_docx_raises(self):
        with pytest.raises(DocumentParsingError, match="Invalid or corrupted"):
            extract_text_from_docx(b"not a valid zip file")

    def test_extract_from_txt_utf8(self, sample_txt_bytes):
        text = extract_text_from_txt(sample_txt_bytes)
        assert "Aisha Khan" in text
        assert "MongoDB" in text

    def test_extract_from_txt_latin1(self):
        content = "Renée Müller\nSkills: Python, C++, Qt".encode("latin-1")
        text = extract_text_from_txt(content)
        assert "Renée Müller" in text
        assert "Python" in text

    def test_extract_from_empty_txt_raises(self):
        with pytest.raises(DocumentParsingError, match="empty"):
            extract_text_from_txt(b"   \n  ")


class TestUnifiedDocumentParser:
    """Test dispatching to appropriate parser based on file extension."""

    def test_dispatch_pdf(self, sample_pdf_bytes):
        text = extract_text_from_document(sample_pdf_bytes, filename="resume.pdf")
        assert "Rahul Sharma" in text
        assert "React" in text

    def test_dispatch_docx(self, sample_docx_bytes):
        text = extract_text_from_document(sample_docx_bytes, filename="resume.docx")
        assert "Marcus Vance" in text

    def test_dispatch_txt(self, sample_txt_bytes):
        text = extract_text_from_document(sample_txt_bytes, filename="resume.txt")
        assert "Aisha Khan" in text

    def test_dispatch_unsupported_format_raises(self):
        with pytest.raises(DocumentParsingError, match="Unsupported document format"):
            extract_text_from_document(b"data", filename="resume.exe")

    def test_mixed_batch_simulation(self, sample_pdf_bytes, sample_docx_bytes, sample_txt_bytes):
        """Simulate a mixed batch of PDF, DOCX, and TXT resumes."""
        batch = [
            ("cand1.pdf", sample_pdf_bytes),
            ("cand2.docx", sample_docx_bytes),
            ("cand3.txt", sample_txt_bytes),
        ]
        parsed_batch = {}
        for fname, raw_bytes in batch:
            text = extract_text_from_document(raw_bytes, filename=fname)
            assert len(text) > 20
            parsed_batch[fname] = text

        assert len(parsed_batch) == 3
        assert "Rahul Sharma" in parsed_batch["cand1.pdf"]
        assert "Marcus Vance" in parsed_batch["cand2.docx"]
        assert "Aisha Khan" in parsed_batch["cand3.txt"]
